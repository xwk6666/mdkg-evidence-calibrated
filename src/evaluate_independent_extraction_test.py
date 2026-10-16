"""Evaluate frozen candidate extraction on the independent blind Gold set."""

from __future__ import annotations

import argparse
import csv
import json
import random
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from disaster_kg_schema import (
    CORE_ENTITY_TYPES,
    CORE_RELATION_TYPES,
    RELATION_SIGNATURES,
    normalize_entity_name,
    normalize_name,
    validate_extraction,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DIR = ROOT / "outputs" / "llm_independent_test_150"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--annotation",
        type=Path,
        default=DEFAULT_DIR / "independent_extraction_test_150_annotation.xlsx",
    )
    parser.add_argument(
        "--predictions",
        type=Path,
        default=DEFAULT_DIR / "independent_extraction_test_150_predictions.json",
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        default=DEFAULT_DIR / "independent_extraction_test_150_manifest.json",
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_DIR / "evaluation")
    parser.add_argument("--bootstrap-repetitions", type=int, default=5000)
    parser.add_argument("--seed", type=int, default=20260802)
    parser.add_argument("--allow-incomplete", action="store_true")
    return parser.parse_args()


def is_complete_status(value: object) -> bool:
    """Accept both the current workbook label and the legacy spelling."""
    return str(value or "").strip().lower() in {"complete", "completed"}


def load_annotation_rows(path: Path) -> List[Dict[str, object]]:
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload.get("samples") if isinstance(payload, dict) else payload
        if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
            raise ValueError("Gold JSON must contain a list or a top-level 'samples' list.")
        normalized_rows: List[Dict[str, object]] = []
        for row in rows:
            normalized = dict(row)
            normalized.setdefault("annotation_status", "complete")
            normalized_rows.append(normalized)
        return normalized_rows

    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading .xlsx requires openpyxl in the active Anaconda environment.") from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    if "Blind_Annotation" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a sheet named 'Blind_Annotation'.")
    sheet = workbook["Blind_Annotation"]
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(value or "").strip() for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]


def parse_json_list(value: object, label: str, blind_id: str) -> List[Dict[str, object]]:
    if value is None or str(value).strip() == "":
        raise ValueError(f"{blind_id}: {label} is blank; use [] when there is no item.")
    try:
        parsed = json.loads(str(value))
    except json.JSONDecodeError as exc:
        raise ValueError(f"{blind_id}: invalid JSON in {label}: {exc}") from exc
    if not isinstance(parsed, list) or not all(isinstance(item, dict) for item in parsed):
        raise ValueError(f"{blind_id}: {label} must be a JSON list of objects.")
    return parsed


def entity_key(entity: Dict[str, str]) -> Tuple[str, str]:
    entity_type = entity["type"]
    name = normalize_entity_name(entity["name"], entity_type).lower()
    return name, entity_type


def relation_key(relation: Dict[str, str]) -> Tuple[str, str, str]:
    return (
        normalize_entity_name(relation["head"], "Disaster").lower(),
        relation["relation"],
        normalize_name(relation["tail"]).lower(),
    )


def counts(predicted: set, gold: set) -> Tuple[int, int, int]:
    return len(predicted & gold), len(predicted - gold), len(gold - predicted)


def metrics(tp: int, fp: int, fn: int) -> Dict[str, float | int]:
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {"tp": tp, "fp": fp, "fn": fn, "precision": precision, "recall": recall, "f1": f1}


def quantile(sorted_values: Sequence[float], probability: float) -> float:
    if not sorted_values:
        return 0.0
    position = (len(sorted_values) - 1) * probability
    lower = int(position)
    upper = min(lower + 1, len(sorted_values) - 1)
    fraction = position - lower
    return sorted_values[lower] * (1 - fraction) + sorted_values[upper] * fraction


def bootstrap_f1(
    sample_counts: Sequence[Tuple[int, int, int]], repetitions: int, rng: random.Random
) -> Dict[str, float | int]:
    if not sample_counts or repetitions <= 0:
        return {"repetitions": 0, "ci95_low": 0.0, "ci95_high": 0.0}
    values: List[float] = []
    n = len(sample_counts)
    for _ in range(repetitions):
        tp = fp = fn = 0
        for _ in range(n):
            sample_tp, sample_fp, sample_fn = sample_counts[rng.randrange(n)]
            tp += sample_tp
            fp += sample_fp
            fn += sample_fn
        values.append(float(metrics(tp, fp, fn)["f1"]))
    values.sort()
    return {
        "repetitions": repetitions,
        "ci95_low": quantile(values, 0.025),
        "ci95_high": quantile(values, 0.975),
    }


def validate_gold_endpoints(
    entities: Sequence[Dict[str, str]], relations: Sequence[Dict[str, object]], blind_id: str
) -> None:
    entity_map = {
        normalize_entity_name(item["name"], item["type"]).lower(): item["type"]
        for item in entities
    }
    for relation in relations:
        relation_type = str(relation.get("relation") or "").strip().upper()
        if relation_type not in RELATION_SIGNATURES:
            raise ValueError(f"{blind_id}: invalid relation type {relation_type!r}.")
        expected_head, expected_tail = RELATION_SIGNATURES[relation_type]
        head = normalize_entity_name(relation.get("head"), expected_head).lower()
        tail = normalize_entity_name(relation.get("tail"), expected_tail).lower()
        if entity_map.get(head) != expected_head or entity_map.get(tail) != expected_tail:
            raise ValueError(
                f"{blind_id}: relation endpoints must also appear in gold_entities_json "
                f"with signature {expected_head}->{expected_tail}."
            )


def write_csv(path: Path, rows: Iterable[Dict[str, object]], fieldnames: Sequence[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    args = parse_args()
    with args.manifest.open("r", encoding="utf-8") as handle:
        manifest = json.load(handle)
    with args.predictions.open("r", encoding="utf-8") as handle:
        predictions = json.load(handle)
    annotation_rows = load_annotation_rows(args.annotation)

    manifest_ids = [sample["blind_id"] for sample in manifest["samples"]]
    prediction_map = {sample["blind_id"]: sample for sample in predictions["samples"]}
    annotation_map = {str(row.get("blind_id") or "").strip(): row for row in annotation_rows}
    if len(annotation_map) != len(annotation_rows):
        raise ValueError("Duplicate or blank blind_id values exist in the annotation workbook.")
    if set(prediction_map) != set(manifest_ids):
        raise ValueError("Predictions and manifest do not contain the same frozen blind IDs.")

    incomplete = [
        blind_id
        for blind_id in manifest_ids
        if not is_complete_status(
            annotation_map.get(blind_id, {}).get("annotation_status")
        )
    ]
    if incomplete and not args.allow_incomplete:
        preview = ", ".join(incomplete[:10])
        raise ValueError(
            f"{len(incomplete)} samples are not complete ({preview}). "
            "Finish annotation before final evaluation."
        )

    entity_total = Counter(tp=0, fp=0, fn=0)
    relation_total = Counter(tp=0, fp=0, fn=0)
    entity_by_type: Dict[str, Counter] = defaultdict(lambda: Counter(tp=0, fp=0, fn=0))
    relation_by_type: Dict[str, Counter] = defaultdict(lambda: Counter(tp=0, fp=0, fn=0))
    entity_sample_counts: List[Tuple[int, int, int]] = []
    relation_sample_counts: List[Tuple[int, int, int]] = []
    error_rows: List[Dict[str, object]] = []
    exact_entity_samples = 0
    exact_relation_samples = 0
    evaluated = 0

    for blind_id in manifest_ids:
        row = annotation_map.get(blind_id)
        if not row or not is_complete_status(row.get("annotation_status")):
            continue
        raw_gold_entities = parse_json_list(row.get("gold_entities_json"), "gold_entities_json", blind_id)
        raw_gold_relations = parse_json_list(row.get("gold_relations_json"), "gold_relations_json", blind_id)
        gold_entities = validate_extraction({"entities": raw_gold_entities, "relations": []})["entities"]
        if len(gold_entities) != len(raw_gold_entities):
            raise ValueError(f"{blind_id}: one or more gold entities violate the closed schema or duplicate another entity.")
        validate_gold_endpoints(gold_entities, raw_gold_relations, blind_id)
        gold_relations = validate_extraction(
            {"entities": gold_entities, "relations": raw_gold_relations}
        )["relations"]
        if len(gold_relations) != len(raw_gold_relations):
            raise ValueError(f"{blind_id}: one or more gold relations are invalid or duplicated.")

        prediction = prediction_map[blind_id]
        cleaned_prediction = validate_extraction(prediction, infer_missing_relations=False)
        predicted_entity_set = {entity_key(item) for item in cleaned_prediction["entities"]}
        gold_entity_set = {entity_key(item) for item in gold_entities}
        predicted_relation_set = {relation_key(item) for item in cleaned_prediction["relations"]}
        gold_relation_set = {relation_key(item) for item in gold_relations}

        entity_counts = counts(predicted_entity_set, gold_entity_set)
        relation_counts = counts(predicted_relation_set, gold_relation_set)
        entity_sample_counts.append(entity_counts)
        relation_sample_counts.append(relation_counts)
        for key, value in zip(("tp", "fp", "fn"), entity_counts):
            entity_total[key] += value
        for key, value in zip(("tp", "fp", "fn"), relation_counts):
            relation_total[key] += value
        exact_entity_samples += int(predicted_entity_set == gold_entity_set)
        exact_relation_samples += int(predicted_relation_set == gold_relation_set)
        evaluated += 1

        for entity_type in CORE_ENTITY_TYPES:
            pred_subset = {item for item in predicted_entity_set if item[1] == entity_type}
            gold_subset = {item for item in gold_entity_set if item[1] == entity_type}
            for key, value in zip(("tp", "fp", "fn"), counts(pred_subset, gold_subset)):
                entity_by_type[entity_type][key] += value
        for relation_type in CORE_RELATION_TYPES:
            pred_subset = {item for item in predicted_relation_set if item[1] == relation_type}
            gold_subset = {item for item in gold_relation_set if item[1] == relation_type}
            for key, value in zip(("tp", "fp", "fn"), counts(pred_subset, gold_subset)):
                relation_by_type[relation_type][key] += value

        manifest_sample = next(item for item in manifest["samples"] if item["blind_id"] == blind_id)
        error_rows.append(
            {
                "blind_id": blind_id,
                "sample_id": manifest_sample["sample_id"],
                "event_family": manifest_sample["event_family"],
                "text": manifest_sample["text"],
                "pred_entities_json": json.dumps(cleaned_prediction["entities"], ensure_ascii=False),
                "gold_entities_json": json.dumps(gold_entities, ensure_ascii=False),
                "pred_relations_json": json.dumps(cleaned_prediction["relations"], ensure_ascii=False),
                "gold_relations_json": json.dumps(gold_relations, ensure_ascii=False),
                "entity_tp": entity_counts[0],
                "entity_fp": entity_counts[1],
                "entity_fn": entity_counts[2],
                "relation_tp": relation_counts[0],
                "relation_fp": relation_counts[1],
                "relation_fn": relation_counts[2],
            }
        )

    if evaluated == 0:
        raise ValueError("No completed samples are available for evaluation.")

    rng = random.Random(args.seed)
    entity_summary = metrics(entity_total["tp"], entity_total["fp"], entity_total["fn"])
    relation_summary = metrics(relation_total["tp"], relation_total["fp"], relation_total["fn"])
    entity_summary.update(bootstrap_f1(entity_sample_counts, args.bootstrap_repetitions, rng))
    relation_summary.update(bootstrap_f1(relation_sample_counts, args.bootstrap_repetitions, rng))
    entity_summary["sample_exact_match"] = exact_entity_samples / evaluated
    relation_summary["sample_exact_match"] = exact_relation_samples / evaluated

    type_rows: List[Dict[str, object]] = []
    for item_type in CORE_ENTITY_TYPES:
        item_metrics = metrics(**entity_by_type[item_type])
        type_rows.append({"task": "entity", "type": item_type, **item_metrics})
    for item_type in CORE_RELATION_TYPES:
        item_metrics = metrics(**relation_by_type[item_type])
        type_rows.append({"task": "relation", "type": item_type, **item_metrics})

    args.output_dir.mkdir(parents=True, exist_ok=True)
    summary = {
        "protocol": manifest["protocol"],
        "evaluated_samples": evaluated,
        "incomplete_samples_skipped": len(incomplete),
        "strict_matching": "normalized entity span + type; normalized head + relation + tail",
        "entity_micro": entity_summary,
        "relation_micro": relation_summary,
        "per_type": type_rows,
    }
    with (args.output_dir / "independent_extraction_metrics.json").open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=False, indent=2)
    write_csv(
        args.output_dir / "independent_extraction_metrics_by_type.csv",
        type_rows,
        ("task", "type", "tp", "fp", "fn", "precision", "recall", "f1"),
    )
    write_csv(
        args.output_dir / "independent_extraction_sample_errors.csv",
        error_rows,
        (
            "blind_id",
            "sample_id",
            "event_family",
            "text",
            "pred_entities_json",
            "gold_entities_json",
            "pred_relations_json",
            "gold_relations_json",
            "entity_tp",
            "entity_fp",
            "entity_fn",
            "relation_tp",
            "relation_fp",
            "relation_fn",
        ),
    )

    print("中文总结")
    print(f"怎么做：在 {evaluated} 条冻结独立样本上进行严格实体和关系匹配。")
    print("原理：模型、提示词与阈值在揭示 Gold 前冻结，评测期间不读取 Gold 且不调参。")
    print(
        "实体：P={precision:.4f} R={recall:.4f} F1={f1:.4f} "
        "95%CI=[{ci95_low:.4f}, {ci95_high:.4f}]".format(**entity_summary)
    )
    print(
        "关系：P={precision:.4f} R={recall:.4f} F1={f1:.4f} "
        "95%CI=[{ci95_low:.4f}, {ci95_high:.4f}]".format(**relation_summary)
    )
    print(f"Artifacts saved to: {args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
